"""End-to-end API tests against a real PostgreSQL database.

Run with:  TEST_DATABASE_URL=postgresql://... pytest
The database is wiped (public schema dropped) at the start of the run.
"""
import os
from io import BytesIO

import psycopg
import pytest

TEST_DB = os.getenv("TEST_DATABASE_URL")
if not TEST_DB:
    pytest.skip("TEST_DATABASE_URL is not set", allow_module_level=True)

os.environ.update(DATABASE_URL=TEST_DB, COOKIE_SECURE="false", ADMIN_USERNAME="admin",
                  ADMIN_PASSWORD="1234", SECRET_KEY="test-secret-key-0123456789")

with psycopg.connect(TEST_DB, autocommit=True) as _c:
    _c.execute("DROP SCHEMA public CASCADE; CREATE SCHEMA public;")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app import main  # noqa: E402


@pytest.fixture(scope="module")
def client():
    with TestClient(main.app) as c:
        yield c


def clear_limits():
    with psycopg.connect(TEST_DB, autocommit=True) as c:
        c.execute("DELETE FROM rate_limit_hits")


@pytest.fixture(autouse=True)
def _reset_limits(client):
    clear_limits()


def test_submit_rate_limit(client):
    bad = {"customer": {"name": "x"}}
    codes = [client.post("/api/feedback", json=bad).status_code for _ in range(11)]
    assert codes[:10] == [422] * 10 and codes[10] == 429


def good_payload(**over):
    p = {
        "customer": {"name": "Ali Raza", "company": "Raza Traders", "phone": "+92 300 1234567",
                     "email": "ali@example.com"},
        "service_type_id": 1,
        "answers": [
            {"question_code": "pickup_time", "value": 4, "suggestion": "Driver came 10 min early"},
            {"question_code": "delivery_time", "value": 5, "detail": "On time"},
            {"question_code": "goods_condition", "value": "Perfect condition"},
            {"question_code": "overall", "value": 5},
            {"question_code": "nps", "value": 9},
        ],
        "other_suggestions": "Keep it up",
    }
    p.update(over)
    return p


def login(client, password="1234"):
    clear_limits()
    return client.post("/api/admin/login", json={"username": "admin", "password": password})


# ---------------------------------------------------------------- public
def test_form_is_served_from_database(client):
    r = client.get("/api/form")
    assert r.status_code == 200
    form = r.json()
    codes = [q["code"] for s in form["sections"] for q in s["questions"]]
    assert "overall" in codes and "driver_behaviour" in codes and len(codes) == 19
    assert len(form["services"]) == 7
    assert client.get("/").status_code == 200
    assert "Content-Security-Policy" in client.get("/").headers


def test_submit_and_reference(client):
    r = client.post("/api/feedback", json=good_payload())
    assert r.status_code == 201, r.text
    assert r.json()["reference"].startswith("AW-")


@pytest.mark.parametrize("change,msg", [
    ({"answers": [{"question_code": "pickup_time", "value": 4}]}, "Please answer"),
    ({"answers": [{"question_code": "overall", "value": 6}]}, "score from 1 to 5"),
    ({"answers": [{"question_code": "overall", "value": True}]}, "valid"),
    ({"answers": [{"question_code": "overall", "value": 5},
                  {"question_code": "goods_condition", "value": "Exploded"}]}, "listed options"),
    ({"answers": [{"question_code": "overall", "value": 5}, {"question_code": "nope", "value": 1}]}, "Unknown question"),
    ({"answers": [{"question_code": "overall", "value": 5}, {"question_code": "overall", "value": 4}]}, "twice"),
    ({"service_type_id": 999}, "service"),
    ({"customer": {"name": ""}}, "at least 1"),
    ({"customer": {"name": "X", "email": "not-an-email"}}, "valid email"),
    ({"customer": {"name": "X"}, "extra": 1}, "Extra inputs"),
])
def test_submission_validation(client, change, msg):
    r = client.post("/api/feedback", json=good_payload(**change))
    assert r.status_code == 422, r.text
    assert msg.lower() in r.json()["detail"].lower()


def test_same_email_reuses_customer_without_overwriting(client):
    client.post("/api/feedback", json=good_payload(customer={"name": "Someone Else", "email": "ALI@example.com"}))
    with psycopg.connect(TEST_DB) as c:
        rows = c.execute("SELECT name FROM customers WHERE lower(email) = 'ali@example.com'").fetchall()
    assert rows == [("Ali Raza",)]


def test_honeypot_stores_nothing(client):
    with psycopg.connect(TEST_DB) as c:
        before = c.execute("SELECT count(*) FROM feedback_responses").fetchone()[0]
    r = client.post("/api/feedback", json=good_payload(website="http://spam"))
    assert r.status_code == 201
    with psycopg.connect(TEST_DB) as c:
        assert c.execute("SELECT count(*) FROM feedback_responses").fetchone()[0] == before


# ----------------------------------------------------------------- admin
def test_admin_requires_login(client):
    client.cookies.clear()
    for path in ("/api/admin/feedback", "/api/admin/stats", "/api/admin/export.xlsx", "/api/admin/me"):
        assert client.get(path).status_code == 401


def test_wrong_password_then_lockout(client):
    client.cookies.clear()
    clear_limits()
    for _ in range(5):
        assert client.post("/api/admin/login", json={"username": "admin", "password": "nope"}).status_code == 401
    assert client.post("/api/admin/login", json={"username": "admin", "password": "1234"}).status_code == 429
    clear_limits()


def test_password_is_hashed(client):
    with psycopg.connect(TEST_DB) as c:
        h = c.execute("SELECT password_hash FROM admin_users WHERE username='admin'").fetchone()[0]
    assert h.startswith("$2") and "1234" not in h


def test_login_list_search_filter_detail_status(client):
    r = login(client)
    assert r.status_code == 200
    cookie = r.headers["set-cookie"].lower()
    assert "httponly" in cookie and "samesite=strict" in cookie
    assert client.get("/api/admin/me").json()["default_password"] is True

    client.post("/api/feedback", json=good_payload(
        customer={"name": "Sara Khan", "phone": "0321-7654321"}, service_type_id=4,
        answers=[{"question_code": "overall", "value": 2, "suggestion": "Truck smelled of fish"},
                 {"question_code": "vehicle_hygiene", "value": 1}]))

    all_ = client.get("/api/admin/feedback").json()
    assert all_["total"] >= 3
    found = client.get("/api/admin/feedback", params={"q": "fish"}).json()
    assert found["total"] == 1 and found["items"][0]["customer_name"] == "Sara Khan"
    assert client.get("/api/admin/feedback", params={"q": "100%_"}).json()["total"] == 0
    low = client.get("/api/admin/feedback", params={"rating_max": 2}).json()
    assert {i["customer_name"] for i in low["items"]} == {"Sara Khan"}
    svc = client.get("/api/admin/feedback", params={"service_id": 4}).json()
    assert svc["total"] == 1
    assert client.get("/api/admin/feedback", params={"rating_min": 9}).status_code == 422
    assert client.get("/api/admin/feedback", params={"date_from": "2000-01-01", "date_to": "2000-01-02"}).json()["total"] == 0

    rid = found["items"][0]["id"]
    detail = client.get(f"/api/admin/feedback/{rid}").json()
    hyg = next(a for a in detail["answers"] if a["code"] == "vehicle_hygiene")
    assert hyg["rating"] == 1
    assert client.patch(f"/api/admin/feedback/{rid}", json={"status": "actioned"}).status_code == 200
    assert client.patch(f"/api/admin/feedback/{rid}", json={"status": "deleted"}).status_code == 422
    assert client.get("/api/admin/feedback", params={"status": "actioned"}).json()["total"] == 1

    s = client.get("/api/admin/stats").json()
    assert s["total"] == all_["total"] and s["low_rated"] == 1 and s["nps"] is not None


def test_cross_site_admin_write_blocked(client):
    login(client)
    r = client.post("/api/admin/logout", headers={"Origin": "https://evil.example"})
    assert r.status_code == 403


def test_excel_export(client):
    login(client)
    client.post("/api/feedback", json=good_payload(
        customer={"name": "=HYPERLINK(\"http://x\",\"click\")"},
        other_suggestions="=1+1"))
    r = client.get("/api/admin/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    ws = wb["Responses"]
    headers = [c.value for c in ws[1]]
    assert "Customer Name" in headers and "Other Suggestions" in headers
    assert any("Suggestion" in h for h in headers)
    assert ws.max_row - 1 == client.get("/api/admin/stats").json()["total"]
    names = [row[headers.index("Customer Name")] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert '=HYPERLINK("http://x","click")' in names  # kept as text, not turned into a formula
    for row in ws.iter_rows(min_row=2):
        cell = row[headers.index("Customer Name")]
        if isinstance(cell.value, str) and cell.value.startswith("="):
            assert cell.data_type == "s"
    assert "Summary" in wb.sheetnames
    filtered = load_workbook(BytesIO(client.get("/api/admin/export.xlsx", params={"q": "fish"}).content))
    assert filtered["Responses"].max_row == 2


def test_change_password_and_logout(client):
    login(client)
    assert client.post("/api/admin/password", json={"current_password": "wrong", "new_password": "abc12345"}).status_code == 400
    assert client.post("/api/admin/password", json={"current_password": "1234", "new_password": "12345678"}).status_code == 422
    assert client.post("/api/admin/password", json={"current_password": "1234", "new_password": "Truck2026x"}).status_code == 200
    assert client.get("/api/admin/me").json()["default_password"] is False
    client.post("/api/admin/logout")
    client.cookies.clear()
    assert client.get("/api/admin/me").status_code == 401
    assert login(client, "1234").status_code == 401
    assert login(client, "Truck2026x").status_code == 200
