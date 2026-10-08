"""Builds the Excel (.xlsx) export."""
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import settings

ORANGE = "F28A1C"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=ORANGE)
STATUS_LABELS = {"new": "New", "reviewed": "Reviewed", "actioned": "Actioned", "archived": "Archived"}


def build_workbook(rows: list[dict], answers: dict[str, dict], questions: list[dict],
                   stats: dict, filters_label: str) -> bytes:
    tz = ZoneInfo(settings.timezone)
    wb = Workbook()

    # -------------------------------------------------------- Responses
    ws = wb.active
    ws.title = "Responses"
    headers = ["Reference", f"Submitted ({settings.timezone})", "Status", "Customer Name",
               "Company", "Phone", "Email", "Service", "Overall (1-5)", "Recommend (0-10)"]
    q_cols: list[tuple[str, str]] = []  # (code, kind)
    for q in questions:
        label = f"{q['section']} – {q['prompt']}"
        scale = " (1-5)" if q["question_type"] == "stars" else " (0-10)" if q["question_type"] == "nps" else ""
        headers.append(label + scale)
        q_cols.append((q["code"], "value"))
        if q["detail_options"]:
            headers.append(label + " – Detail")
            q_cols.append((q["code"], "detail"))
        headers.append(label + " – Suggestion")
        q_cols.append((q["code"], "suggestion"))
    headers.append("Other Suggestions")
    ws.append(headers)

    text_cols: set[int] = {1, 4, 5, 6, 7, 8, len(headers)}
    for i, (_, kind) in enumerate(q_cols, start=11):
        if kind in ("detail", "suggestion"):
            text_cols.add(i)

    for r in rows:
        given = answers.get(str(r["id"]), {})
        line = [
            r["reference_code"],
            r["submitted_at"].astimezone(tz).replace(tzinfo=None),
            STATUS_LABELS.get(r["status"], r["status"]),
            r["customer_name"], r["customer_company"], r["customer_phone"], r["customer_email"],
            r["service"], r["overall_rating"], r["nps_score"],
        ]
        for code, kind in q_cols:
            a = given.get(code)
            if a is None:
                line.append(None)
            elif kind == "value":
                line.append(a["rating"] if a["rating"] is not None else a["choice_value"])
            elif kind == "detail":
                line.append(a["detail_value"])
            else:
                line.append(a["suggestion"])
        line.append(r["other_suggestions"])
        ws.append(line)
        # openpyxl treats any string starting with "=" as a formula; force typed text to stay text.
        row_idx = ws.max_row
        for col in text_cols:
            cell = ws.cell(row=row_idx, column=col)
            if isinstance(cell.value, str):
                cell.data_type = "s"

    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    widths = [16, 18, 11, 22, 20, 16, 26, 22, 12, 14]
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1] if i <= len(widths) else 24
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "yyyy-mm-dd hh:mm"

    # ---------------------------------------------------------- Summary
    sm = wb.create_sheet("Summary")
    sm.append(["Ashwheelz customer feedback"])
    sm["A1"].font = Font(bold=True, size=14, color=ORANGE)
    sm.append(["Exported", datetime.now(tz).replace(tzinfo=None)])
    sm["B2"].number_format = "yyyy-mm-dd hh:mm"
    sm.append(["Filters", filters_label])
    sm.append([])
    sm.append(["Responses", stats["total"]])
    sm.append(["Average overall rating (1-5)", stats["avg_overall"]])
    sm.append(["Net Promoter Score (-100 to 100)", stats["nps"]])
    sm.append(["Rated 1-2 stars overall", stats["low_rated"]])
    sm.append([])
    sm.append(["Question", "Average (1-5)", "Answers"])
    hdr = sm.max_row
    for q in stats["by_question"]:
        sm.append([f"{q['section']} – {q['prompt']}", q["avg"], q["n"]])
    for cell in sm[hdr]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r in range(5, 9):
        sm.cell(row=r, column=1).font = Font(bold=True)
    sm.column_dimensions["A"].width = 62
    sm.column_dimensions["B"].width = 18
    sm.column_dimensions["C"].width = 10
    sm["B3"].data_type = "s"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
